"""Enterprise controls: scope, audit, approvals, calendar, production guard."""

from datetime import UTC, datetime

import pytest
from httpx import AsyncClient

from app.core.calendar import add_business_minutes, is_business_day
from app.core.config import Settings
from app.core.jalali import gregorian_to_jalali
from app.core.security import totp_now, verify_totp
from app.core.startup import assert_production_safe
from tests.conftest import auth, login


@pytest.mark.asyncio
async def test_branch_scope_hides_other_branch_work(client: AsyncClient, seeded):
    token = await login(client, "ali@autoco.example")
    other = await client.post(
        "/org/branches",
        headers=auth(token),
        json={"name": "Shiraz", "code": "SHZ"},
    )
    # Ali is a manager, not org admin, so branch create is forbidden.
    assert other.status_code == 403

    created = await client.post(
        "/work-items",
        headers=auth(token),
        json={"title": "In Ali branch", "branch_id": str(seeded["branch"].id)},
    )
    assert created.status_code == 200
    listed = await client.get("/work-items", headers=auth(token))
    assert listed.status_code == 200
    assert any(item["title"] == "In Ali branch" for item in listed.json())


@pytest.mark.asyncio
async def test_self_scope_hides_others_work(client: AsyncClient, seeded):
    ali = await login(client, "ali@autoco.example")
    made = await client.post(
        "/work-items",
        headers=auth(ali),
        json={"title": "Ali only"},
    )
    assert made.status_code == 200
    viewer = await login(client, "viewer@autoco.example")
    listed = await client.get("/work-items", headers=auth(viewer))
    assert listed.status_code == 200
    assert all(item["title"] != "Ali only" for item in listed.json())


@pytest.mark.asyncio
async def test_timeline_is_org_scoped(client: AsyncClient, seeded):
    token = await login(client, "ali@autoco.example")
    made = await client.post(
        "/work-items",
        headers=auth(token),
        json={"title": "Timeline item"},
    )
    wid = made.json()["id"]
    timeline = await client.get(f"/timeline/work_item/{wid}", headers=auth(token))
    assert timeline.status_code == 200
    assert timeline.json()


@pytest.mark.asyncio
async def test_login_writes_audit(client: AsyncClient, seeded):
    token = await login(client, "ali@autoco.example")
    audit = await client.get("/audit", headers=auth(token))
    assert audit.status_code == 200
    assert any(row["action"] == "auth.login" for row in audit.json())


@pytest.mark.asyncio
async def test_approval_two_steps(client: AsyncClient, seeded):
    ali = await login(client, "ali@autoco.example")
    reza = await login(client, "reza@autoco.example")
    created = await client.post(
        "/approvals",
        headers={**auth(reza), "Idempotency-Key": "approve-1"},
        json={
            "kind": "discount",
            "title": "تخفیف ۱۰ درصد",
            "payload": {"percent": 10},
            # درخواست‌دهنده (reza) نباید تأییدکننده باشد — خود-تأییدی ممنوع
            "approver_ids": [str(seeded["ali"].id), str(seeded["sara"].id)],
        },
    )
    assert created.status_code == 200, created.text
    again = await client.post(
        "/approvals",
        headers={**auth(reza), "Idempotency-Key": "approve-1"},
        json={
            "kind": "discount",
            "title": "تخفیف ۱۰ درصد",
            "approver_ids": [str(seeded["ali"].id), str(seeded["sara"].id)],
        },
    )
    assert again.json()["id"] == created.json()["id"]
    request_id = created.json()["id"]
    first = await client.post(
        f"/approvals/{request_id}/decision",
        headers=auth(ali),
        json={"decision": "approved"},
    )
    assert first.json()["status"] == "pending"
    sara = await login(client, "sara@autoco.example")
    second = await client.post(
        f"/approvals/{request_id}/decision",
        headers=auth(sara),
        json={"decision": "approved"},
    )
    assert second.json()["status"] == "approved"


@pytest.mark.asyncio
async def test_kpi_export(client: AsyncClient, seeded):
    token = await login(client, "ali@autoco.example")
    kpis = await client.get("/reports/kpis", headers=auth(token))
    assert kpis.status_code == 200
    assert "as_of_jalali" in kpis.json()
    csv_resp = await client.get("/reports/kpis.csv", headers=auth(token))
    assert csv_resp.status_code == 200
    assert "open_tasks" in csv_resp.text


@pytest.mark.asyncio
async def test_kpi_xlsx_export(client: AsyncClient, seeded):
    import io

    from openpyxl import load_workbook

    token = await login(client, "ali@autoco.example")
    resp = await client.get("/reports/kpis.xlsx", headers=auth(token))
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert "open_tasks" in rows
    assert "as_of_jalali" in rows

    audit = await client.get("/audit?action=report.export", headers=auth(token))
    assert audit.status_code == 200
    formats = [row for row in audit.json() if row["action"] == "report.export"]
    assert formats, "expected report.export audit row"


@pytest.mark.asyncio
async def test_kpi_pdf_export(client: AsyncClient, seeded):
    token = await login(client, "ali@autoco.example")
    resp = await client.get("/reports/kpis.pdf", headers=auth(token))
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")
    assert b"open_tasks" in resp.content
    assert b"%%EOF" in resp.content


@pytest.mark.asyncio
async def test_export_denied_without_event_read(client: AsyncClient, seeded):
    """Viewer has WORK_ITEM_READ only — every KPI export must 403."""
    token = await login(client, "viewer@autoco.example")
    for path in ("/reports/kpis.xlsx", "/reports/kpis.pdf"):
        resp = await client.get(path, headers=auth(token))
        assert resp.status_code == 403, path


@pytest.mark.asyncio
async def test_v1_prefix(client: AsyncClient, seeded):
    token = await login(client, "ali@autoco.example")
    health = await client.get("/health")
    assert health.status_code == 200
    ready = await client.get("/health/ready")
    assert ready.status_code == 200
    metrics = await client.get("/metrics")
    assert "cm_http_requests_total" in metrics.text
    people = await client.get("/v1/people", headers=auth(token))
    assert people.status_code == 200


def test_jalali_and_business_day():
    assert gregorian_to_jalali(2026, 3, 21)[0] == 1405
    holiday = datetime(2026, 3, 20, 10, 0, tzinfo=UTC)
    assert is_business_day(holiday.date()) is False
    moved = add_business_minutes(holiday, 30)
    assert moved.date() != holiday.date()


def test_totp_roundtrip():
    secret = "cm-test-secret"
    code = totp_now(secret)
    assert verify_totp(secret, code)


def test_production_refuses_placeholder_secret():
    settings = Settings(
        app_env="production",
        debug=False,
        secret_key="dev-secret-change-me",
        database_url="postgresql+asyncpg://cm:cm@localhost/cm",
    )
    with pytest.raises(RuntimeError):
        assert_production_safe(settings)
