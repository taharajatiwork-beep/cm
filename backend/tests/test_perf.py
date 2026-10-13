"""پرفورمنس: قوانین پیامک، لیست، پرداخت، CPL، اکسل و مجوز."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.core.permissions import PermissionCode
from app.perf.services.excel import COST_HEADERS, costs_xlsx
from app.perf.services.lists import clean_numbers, panel_number
from app.perf.sms_rules import check_sms, has_opt_out, page_count
from app.services import OrgService
from tests.conftest import auth, login


def test_ledger_cost_replaces_loose_rows():
    from types import SimpleNamespace
    from app.perf.services.reports import report_from_parts

    campaign = SimpleNamespace(id="c", code="FM54", title="فروش سازمانی", brand="فونیکس و مدیران", starts_on=None, ends_on=None, budget_toman=None)
    send = SimpleNamespace(cost_toman=999, leads=1, sent_at=None, calls=0, views=0)
    cost = SimpleNamespace(campaign_code="FM54", amount_toman=4_000_000, kind="پیامک", counts_for_cpl=True, sms_send_id=None, story_id=None, spent_on=None)
    report = report_from_parts(campaign, [send], [], [], {}, None, None, [cost])
    assert report["cost_toman"] == 4_000_000
    assert report["sms_cost_toman"] == 4_000_000


def test_next_step_follows_the_open_gap_and_stops_before_spending():
    from types import SimpleNamespace
    from app.perf.services.queue import campaign_gap, next_steps

    campaign = SimpleNamespace(id="c1", code="S84", status="active", budget_toman=0, channels="sms,ig_story")
    car = SimpleNamespace(campaign_id="c1", status="new")
    budget_gap = campaign_gap(campaign, [car], [], [])
    assert budget_gap.get("label") == "بودجه"
    assert budget_gap["href"].endswith("step=budget")
    campaign.budget_toman = 80_000_000
    assert campaign_gap(campaign, [car], [], []).get("label") == "چیدن پیج"
    story = SimpleNamespace(campaign_id="c1", status="planned", views=0, impressions=0, leads=0)
    assert campaign_gap(campaign, [car], [story], []).get("label") == "انتشار"
    story.status = "published"
    assert campaign_gap(campaign, [car], [story], []).get("label") == "اینسایت"
    story.views = 10
    send = SimpleNamespace(campaign_id="c1", recipient_count=20, delivered_count=0, sent_at=None)
    assert campaign_gap(campaign, [car], [story], []).get("label") == "پیامک"
    send.sent_at = "2026-09-01"
    assert campaign_gap(campaign, [car], [story], [send]).get("label") == "تحویل"
    done = SimpleNamespace(id="c2", code="S1", status="done", budget_toman=0, channels="sms")
    panel = {"name": "آی‌اس‌ام‌اس", "low": True}
    rows = next_steps([done, campaign], [panel], 2, [car], [story], [send])
    assert [row["label"] for row in rows[:3]] == ["شارژ", "پول", "تحویل"]
    assert rows[2]["title"].startswith("S84")
    assert all("S1" not in row["title"] for row in rows)


def test_next_campaign_code_follows_the_register():
    from app.perf.campaign_register import book_max, next_code

    assert book_max() == 93
    assert next_code("U", ["S84", "S90"]) == "U94"
    assert next_code("ni", ["AC100"]) == "NI101"


def test_extract_channel_codes_keep_sms_and_social_apart():
    from app.perf.services.extract_import import channel_codes, toman_from_rial

    assert channel_codes(["پیامک دیتابیس", "اینستاگرام", "دیوار"]) == ["sms", "ig_story", "divar"]
    assert toman_from_rial(17_400_962) == 1_740_096


def test_page_count_and_opt_out():
    assert page_count("س" * 70) == 1
    assert page_count("س" * 71) == 2
    assert page_count("س" * 134) == 2
    assert page_count("س" * 135) == 3
    assert page_count("a" * 160) == 1
    assert page_count("a" * 161) == 2
    assert has_opt_out("سلام\nلغو11")
    assert has_opt_out("سلام لغو۱۱")
    checked = check_sms("فونیکس تماس 02147677777 https://x.ir\nلغو11", brand="فونیکس")
    assert checked["ok"]
    assert checked["pages"] == 1
    long_text = "س" * 700 + "لغو11"
    assert page_count(long_text) > 10


def test_drafts_keep_one_page_under_seventy():
    import asyncio

    from app.perf.services.sms import draft_messages
    from app.perf.sms_rules import billed_len

    data = asyncio.run(
        draft_messages(
            brand="فونیکس",
            cars="تیگو ۸ پرو مکس " * 6,
            offer="اقساط بدون بهره و تحویل فوری در شعب مرزداران و ایرانمال",
            phone="021-47677777",
            link="https://www.sotoodeh.com/cars?utm=campaign",
            campaign_code="S84",
        )
    )
    by_kind = {item["kind"]: item for item in data["messages"]}
    ad1 = by_kind["ad1"]
    assert ad1["check"]["ok"]
    assert ad1["check"]["pages"] == 1
    assert billed_len(ad1["text"]) <= 70
    assert ad1["text"].endswith("لغو11")
    assert "S84" not in ad1["text"]
    assert "فونیکس" in ad1["text"]
    assert by_kind["ad2"]["check"]["pages"] <= 2
    assert billed_len(by_kind["ad2"]["text"]) <= 134
    assert by_kind["service"]["check"]["pages"] == 1
    assert by_kind["service"]["text"].startswith("مشتری عزیز")
    assert "متقاضی گرامی" not in by_kind["service"]["text"]
    assert "S84" not in by_kind["service"]["text"]
    short = asyncio.run(
        draft_messages(
            brand="هیوندا",
            cars="توسان",
            offer="تحویل فوری",
            phone="02147677777",
            link="",
            campaign_code="S84",
        )
    )
    assert short["messages"][0]["text"] == "هیوندا\nتوسان، تحویل فوری\n021-47677777\nستوده\nلغو11"
    assert "S84" not in short["messages"][2]["text"]


def test_list_cleanup_normalizes_and_drops():
    cleaned = clean_numbers(
        ["09121111111", "989121111111", "9121111111", "09122222222", "09133333333", "نه"],
        blacklist={"989122222222"},
        recent={"989133333333"},
    )
    assert cleaned["duplicate_count"] == 2
    assert cleaned["blacklist_count"] == 1
    assert cleaned["recent_count"] == 1
    assert cleaned["invalid_count"] == 1
    assert cleaned["phones"] == ["989121111111"]
    assert panel_number("989121111111") == "9121111111"


def test_sheet_report_comes_from_calculated_tab():
    from types import SimpleNamespace
    from uuid import uuid4

    from app.perf.services.sheets import apply_report_values, build_tables, spreadsheet_id

    assert spreadsheet_id("https://docs.google.com/spreadsheets/d/abc123/edit?usp=sharing") == "abc123"
    campaign = SimpleNamespace(id=uuid4(), code="S84", title="فونیکس", brand="فونیکس")
    tables = build_tables([campaign], [], [], [], [], {})
    assert tables["Report"][1][0] == "S84"
    assert tables["Report"][1][3].startswith("=SUMIF(SMS!")
    assert "Social!E:E" in tables["Report"][1][3]
    reports = {campaign.id: {"cost_toman": 1, "views": 1, "leads": 1, "cpl": 1, "cpv": 1, "source": "sheet"}}
    apply_report_values([campaign], reports, [["کمپین"], ["S84", "فونیکس", "فونیکس", 9400000, 120000, 40, 235000, 78.3]])
    assert reports[campaign.id]["cost_toman"] == 9400000
    assert reports[campaign.id]["leads"] == 40
    assert reports[campaign.id]["cpl"] == 235000
    assert reports[campaign.id]["source"] == "google"


def test_cost_workbook_headers():
    data = costs_xlsx(
        [
            {
                "cost_number": 1,
                "spent_on": "2026-01-01",
                "amount_toman": 1000,
                "kind": "پیامک",
                "detail": "",
                "brand": "فونیکس",
                "campaign_code": "S1",
                "cost_center": "پرفورمنس",
                "description": "تست",
            }
        ]
    )
    book = load_workbook(BytesIO(data))
    assert [cell.value for cell in book.active[1]] == COST_HEADERS


async def _nadia(db, seeded):
    svc = OrgService(db)
    role = await svc.create_role(
        seeded["org"].id,
        "پرفورمنس",
        "performance_marketing",
        [
            PermissionCode.PERF_READ,
            PermissionCode.PERF_WRITE,
            PermissionCode.PERF_APPROVE,
            PermissionCode.PRICE_APPROVE,
            PermissionCode.NOTIFICATION_READ,
        ],
        "organization",
    )
    user = await svc.create_user(
        organization_id=seeded["org"].id,
        email="nadia-test@autoco.example",
        full_name="نادیا تست",
        password="password",
        role_id=role.id,
        branch_id=seeded["branch"].id,
    )
    await db.commit()
    return user


def _xlsx(rows: list[list]) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


async def test_seller_cannot_open_perf(client, seeded):
    token = await login(client, "reza@autoco.example")
    response = await client.get("/perf/campaigns", headers=auth(token))
    assert response.status_code == 403


async def test_payment_records_cost_and_campaign_cpl(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)

    panel = await client.post(
        "/perf/panels",
        headers=headers,
        json={"code": "isms", "name": "آی‌اس‌ام‌اس", "sender_line": "5000", "tariff_toman": 100, "balance_toman": 1000, "low_balance_toman": 100},
    )
    assert panel.status_code == 201, panel.text
    panel_id = panel.json()["id"]

    campaign = await client.post(
        "/perf/campaigns",
        headers=headers,
        json={"code": "S84", "title": "فروردین", "brand": "فونیکس", "channels": ["sms"]},
    )
    assert campaign.status_code == 201, campaign.text
    campaign_id = campaign.json()["id"]

    send = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={
            "campaign_id": campaign_id,
            "panel_id": panel_id,
            "body": "فونیکس تماس 02147677777 https://x.ir",
            "recipient_count": 4,
            "cost_toman": 1000,
            "leads": 4,
        },
    )
    assert send.status_code == 201, send.text
    assert send.json()["body"].endswith("لغو11")
    assert send.json()["pages"] >= 1

    detail = await client.get(f"/perf/campaigns/{campaign_id}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["report"]["cost_toman"] == 1000
    assert detail.json()["report"]["leads"] == 4
    assert detail.json()["report"]["cpl"] == 250

    payment = await client.post(
        "/perf/payments",
        headers=headers,
        json={
            "kind": "panel_charge",
            "amount_toman": 500,
            "beneficiary": "پنل",
            "iban": "IR123",
            "panel_id": panel_id,
            "campaign_id": campaign_id,
            "cost_center": "پرفورمنس",
        },
    )
    assert payment.status_code == 201, payment.text
    payment_id = payment.json()["id"]
    own = await client.post(f"/perf/payments/{payment_id}/advance", headers=headers, json={"action": "finance"})
    assert own.status_code == 403
    finance_role = await OrgService(db).create_role(
        seeded["org"].id,
        "مالی تست",
        "finance_test",
        [PermissionCode.PRICE_APPROVE, PermissionCode.PERF_READ, PermissionCode.NOTIFICATION_READ],
        "organization",
    )
    await OrgService(db).create_user(
        organization_id=seeded["org"].id,
        email="finance-test@autoco.example",
        full_name="مالی تست",
        password="password",
        role_id=finance_role.id,
        branch_id=seeded["branch"].id,
    )
    await db.commit()
    finance_token = await login(client, "finance-test@autoco.example")
    stale = await client.post(
        f"/perf/payments/{payment_id}/advance",
        headers=auth(finance_token),
        json={"action": "finance", "seen_amount_toman": 999},
    )
    assert stale.status_code == 409, stale.text
    step = await client.post(
        f"/perf/payments/{payment_id}/advance",
        headers=auth(finance_token),
        json={"action": "finance"},
    )
    assert step.status_code == 200, step.text
    assert step.json()["status"] == "recorded"

    panels = await client.get("/perf/panels", headers=headers)
    balance = next(row["balance_toman"] for row in panels.json() if row["id"] == panel_id)
    assert balance == 500  # 1000 - 1000 send + 500 charge

    exported = await client.get("/perf/excel/costs", headers=headers)
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content))
    assert book.active["A1"].value == "شماره هزینه"
    assert book.active.max_row >= 3

    missing = await client.get(f"/perf/integrations/pishgam/balance?panel_id={panel_id}", headers=headers)
    assert missing.status_code == 400

    sheet = await client.post("/perf/integrations/sheets", headers=headers, json={"kind": "costs"})
    assert sheet.status_code == 200
    assert sheet.json()["pushed"] is False


async def test_prepare_list_from_didar_sheet(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)
    await client.post("/perf/blacklist", headers=headers, json={"phone": "09124444444", "note": "لغو"})
    data = _xlsx(
        [
            ["موبایل"],
            ["09121111111"],
            ["09121111111"],
            ["09124444444"],
            ["نامعتبر"],
        ]
    )
    prepared = await client.post(
        "/perf/lists/prepare",
        headers=headers,
        files={"file": ("didar.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"recent_days": "14", "chunk_size": "15000"},
    )
    assert prepared.status_code == 201, prepared.text
    body = prepared.json()
    assert body["final_count"] == 1
    assert body["duplicate_count"] == 1
    assert body["blacklist_count"] == 1
    exported = await client.get(f"/perf/lists/{body['id']}/export", headers=headers)
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content)).active
    assert sheet["A2"].value == "9121111111"


async def test_page_plan_counts_only_after_payment_and_publish(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)
    campaign = await client.post(
        "/perf/campaigns",
        headers=headers,
        json={"code": "P1", "title": "پیج", "brand": "فونیکس", "channels": ["ig_story"], "budget_toman": 5_000_000},
    )
    assert campaign.status_code == 201, campaign.text
    campaign_id = campaign.json()["id"]
    page = await client.post(
        "/perf/pages",
        headers=headers,
        json={"name": "خودرو بانک", "topic": "خودرو", "contact": "ادمین", "price_toman": 2_000_000},
    )
    assert page.status_code == 201, page.text
    placed = await client.post(
        "/perf/stories",
        headers=headers,
        json={"campaign_id": campaign_id, "page_id": page.json()["id"], "kind": "story"},
    )
    assert placed.status_code == 201, placed.text
    body = placed.json()
    assert body["status"] == "planned"
    assert body["payment_status"] == "requested"
    assert body["cost_toman"] == 2_000_000

    detail = await client.get(f"/perf/campaigns/{campaign_id}", headers=headers)
    assert detail.json()["report"]["cost_toman"] == 0
    assert detail.json()["report"]["pages"] == []

    finance_role = await OrgService(db).create_role(
        seeded["org"].id,
        "مالی پیج",
        "finance_page",
        [PermissionCode.PRICE_APPROVE, PermissionCode.PERF_READ, PermissionCode.NOTIFICATION_READ],
        "organization",
    )
    await OrgService(db).create_user(
        organization_id=seeded["org"].id,
        email="finance-page@autoco.example",
        full_name="مالی پیج",
        password="password",
        role_id=finance_role.id,
        branch_id=seeded["branch"].id,
    )
    await db.commit()
    finance = auth(await login(client, "finance-page@autoco.example"))
    paid = await client.post(
        f"/perf/payments/{body['payment_id']}/advance",
        headers=finance,
        json={"action": "finance"},
    )
    assert paid.status_code == 200, paid.text
    still = await client.get(f"/perf/campaigns/{campaign_id}", headers=headers)
    assert still.json()["report"]["cost_toman"] == 0

    published = await client.post(
        f"/perf/stories/{body['id']}/publish",
        headers=headers,
        json={"published_on": "2026-09-12", "views": 1000, "leads": 4},
    )
    assert published.status_code == 200, published.text
    assert published.json()["status"] == "published"
    done = await client.get(f"/perf/campaigns/{campaign_id}", headers=headers)
    report = done.json()["report"]
    assert report["cost_toman"] == 2_000_000
    assert report["leads"] == 4
    assert report["pages"][0]["name"] == "خودرو بانک"
    assert report["pages"][0]["cpl"] == 500_000

    changed = await client.patch(
        f"/perf/pages/{page.json()['id']}",
        headers=headers,
        json={"price_toman": 2_500_000},
    )
    assert changed.status_code == 200, changed.text
    listed = await client.get("/perf/pages", headers=headers)
    row = next(item for item in listed.json() if item["id"] == page.json()["id"])
    assert row["price_toman"] == 2_500_000
    assert row["prices"][1]["price_toman"] == 2_000_000


async def test_short_balance_opens_an_editable_charge(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)
    panel = await client.post(
        "/perf/panels",
        headers=headers,
        json={"code": "isms", "name": "کم", "tariff_toman": 100, "balance_toman": 1000, "low_balance_toman": 100},
    )
    assert panel.status_code == 201, panel.text
    panel_id = panel.json()["id"]
    blocked = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={"panel_id": panel_id, "body": "سلام", "recipient_count": 10, "cost_toman": 5000},
    )
    assert blocked.status_code == 409, blocked.text
    assert "قابل ویرایش" in blocked.json()["detail"]
    payments = await client.get("/perf/payments", headers=headers)
    charge = next(row for row in payments.json() if row["panel_id"] == panel_id)
    assert charge["amount_toman"] == 4000
    assert charge["status"] == "requested"
    edited = await client.patch(f"/perf/payments/{charge['id']}", headers=headers, json={"amount_toman": 8000})
    assert edited.status_code == 200, edited.text
    assert edited.json()["amount_toman"] == 8000
    again = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={"panel_id": panel_id, "body": "سلام دوباره", "recipient_count": 10, "cost_toman": 5000},
    )
    assert again.status_code == 409
    listed = await client.get("/perf/payments", headers=headers)
    assert sum(1 for row in listed.json() if row["panel_id"] == panel_id and row["status"] == "requested") == 1


async def test_delivery_file_matches_the_send(client, seeded, db):
    await _nadia(db, seeded)
    headers = auth(await login(client, "nadia-test@autoco.example"))
    panel = await client.post(
        "/perf/panels",
        headers=headers,
        json={"code": "isms", "name": "تحویل", "balance_toman": 100_000, "tariff_toman": 100},
    )
    campaign = await client.post(
        "/perf/campaigns",
        headers=headers,
        json={"code": "D1", "title": "تحویل", "brand": "فونیکس", "channels": ["sms"]},
    )
    send = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={
            "campaign_id": campaign.json()["id"],
            "panel_id": panel.json()["id"],
            "body": "سلام",
            "recipient_count": 3,
            "cost_toman": 300,
        },
    )
    assert send.status_code == 201, send.text
    # کد کمپین روی ارسال می‌ماند، نه داخل متن پیامک (docs/performance-marketing.md)
    assert send.json()["campaign_code"] == "D1"
    assert "D1" not in send.json()["body"]
    book = Workbook()
    sheet = book.active
    sheet.append(["وضعیت"])
    sheet.append(["دریافت موفق"])
    sheet.append(["ارسال شده به مخابرات منتظر دریافت گزارش"])
    sheet.append(["دریافت موفق"])
    buf = BytesIO()
    book.save(buf)
    matched = await client.post(
        "/perf/sms/delivery",
        headers=headers,
        files={"file": ("report.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert matched.status_code == 200, matched.text
    assert matched.json()["delivered_count"] == 2
    assert matched.json()["campaign_code"] == "D1"


async def test_insight_xlsx_fills_numbers_without_llm(client, seeded, db):
    await _nadia(db, seeded)
    headers = auth(await login(client, "nadia-test@autoco.example"))
    campaign = await client.post(
        "/perf/campaigns",
        headers=headers,
        json={"code": "I1", "title": "اینسایت", "brand": "فونیکس", "channels": ["ig_story"]},
    )
    page = await client.post("/perf/pages", headers=headers, json={"name": "پیج فایل", "price_toman": 0})
    placed = await client.post(
        "/perf/stories",
        headers=headers,
        json={"campaign_id": campaign.json()["id"], "page_id": page.json()["id"], "kind": "story"},
    )
    published = await client.post(
        f"/perf/stories/{placed.json()['id']}/publish",
        headers=headers,
        json={"published_on": "2026-09-01"},
    )
    assert published.status_code == 200, published.text
    book = Workbook()
    sheet = book.active
    sheet.append(["نمایش", "ورودی"])
    sheet.append([1200, 6])
    buf = BytesIO()
    book.save(buf)
    read = await client.post(
        f"/perf/stories/{placed.json()['id']}/insight",
        headers=headers,
        files={"file": ("insight.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert read.status_code == 200, read.text
    assert read.json()["ok"] is True
    assert read.json()["fields"]["views"] == 1200
    assert read.json()["fields"]["leads"] == 6


async def test_edit_and_delete_campaign_removes_its_send(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)
    campaign = await client.post(
        "/perf/campaigns",
        headers=headers,
        json={"code": "Z9", "title": "موقت", "brand": "تست", "channels": ["sms"]},
    )
    assert campaign.status_code == 201, campaign.text
    campaign_id = campaign.json()["id"]
    renamed = await client.patch(f"/perf/campaigns/{campaign_id}", headers=headers, json={"title": "موقت ۲"})
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["title"] == "موقت ۲"
    send = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={"campaign_id": campaign_id, "body": "تست حذف", "recipient_count": 1, "cost_toman": 10},
    )
    assert send.status_code == 201, send.text
    gone = await client.delete(f"/perf/campaigns/{campaign_id}", headers=headers)
    assert gone.status_code == 204, gone.text
    missing = await client.get(f"/perf/campaigns/{campaign_id}", headers=headers)
    assert missing.status_code == 404


def test_source_sheets_append_only_the_new_row():
    from datetime import date
    from types import SimpleNamespace

    from app.perf.services.source_sheets import cost_rows, describe_sources, register_rows, social_tab

    old = SimpleNamespace(code="S93", title="قبلی", starts_on=date(2026, 9, 1), created_at=None)
    fresh = SimpleNamespace(code="S95", title="تست شیت", starts_on=date(2026, 10, 5), created_at=None)
    legacy = SimpleNamespace(code="K1-1", title="قدیمی", starts_on=None, created_at=None)
    rows = register_rows([old, fresh, legacy], {"S93"})
    assert rows == [["S", 95, "S95", "تست شیت", "1405-07-13"]]

    cost = SimpleNamespace(
        cost_number=988,
        spent_on=date(2026, 10, 5),
        amount_toman=333000,
        kind="تست",
        detail="تست شیت",
        brand="ترکیبی",
        campaign_code="S95",
        cost_center="پرفورمنس",
        description="تست شیت هزینه",
    )
    duplicate = SimpleNamespace(
        cost_number=1,
        spent_on=date(2026, 1, 1),
        amount_toman=1050000,
        kind="چاپی",
        detail="",
        brand="فونیکس",
        campaign_code="",
        cost_center="",
        description="سفارش ست اداری چاپی",
    )
    sheet, log = cost_rows([duplicate, cost], {("", 1050000, "سفارش ست اداری چاپی")})
    assert len(sheet) == 1 and sheet[0][0] == 988 and sheet[0][6] == "S95"
    assert log[0][3] == 333000
    assert social_tab("خودرو بانک", {"خودرو بانک"}, set()) == "BaleRubika"
    assert social_tab("ستوده مگ", set(), {"ستوده مگ"}) == "Instagram-S"
    assert describe_sources({"register": 1, "costs": 2, "sms": 0, "social": 1}) == (
        "به شیت‌های اصلی اضافه شد: دفتر کمپین 1، هزینه 2، سوشال 1."
    )
    from datetime import datetime as stamp

    from app.perf.services.source_sheets import pending_source_export

    sent = SimpleNamespace(source_sheet_exported_at=stamp(2026, 10, 5, 8, 0))
    pending = SimpleNamespace(source_sheet_exported_at=None)
    assert pending_source_export([sent, pending]) == [pending]


def test_mizito_reader_keeps_open_marketing_cards():
    from app.perf.services.mizito import campaign_code, read_cards, wanted_project

    assert wanted_project("پرفورمنس مارکتینگ")
    assert not wanted_project("سئو")
    assert campaign_code("پرداخت استوری S95") == "S95"
    parsed = read_cards(
        {"projects": [
            {"_id": "p1", "title": "مالی و مارکتینگ", "kanban_boards": [{"_id": "b1", "title": "درخواست پرداخت"}, {"_id": "b2", "title": "کنسل شد"}]},
            {"_id": "p2", "title": "سئو", "kanban_boards": []},
        ]},
        {"tasks": [
            {"title": "شارژ پنل S95", "project": "p1", "kanban_board": "b1", "completed": False},
            {"title": "تمام‌شده", "project": "p1", "kanban_board": "b2", "completed": False},
            {"title": "پرداخت‌شده", "project": "p1", "kanban_board": "b1", "completed": True},
            {"title": "بلاگ", "project": "p2", "kanban_board": "", "completed": False},
        ]},
    )
    assert parsed["open_count"] == 1
    assert parsed["cards"][0]["board"] == "درخواست پرداخت"
    assert parsed["cards"][0]["code"] == "S95"


async def test_saved_mizito_list_is_what_the_page_reads(client, seeded, db):
    import json

    from app.perf.models import PmIntegration

    await _nadia(db, seeded)
    db.add(
        PmIntegration(
            organization_id=seeded["org"].id,
            mizito_token="secret-token-value",
            mizito_snapshot=json.dumps(
                {
                    "connected": True,
                    "reason": "",
                    "projects": [{"title": "مالی و مارکتینگ", "id": "p1"}],
                    "cards": [{"project": "مالی و مارکتینگ", "board": "درخواست پرداخت", "title": "شارژ S95", "code": "S95"}],
                    "open_count": 1,
                },
                ensure_ascii=False,
            ),
        )
    )
    await db.commit()
    token = await login(client, "nadia-test@autoco.example")
    page = await client.get("/perf/integrations/mizito", headers=auth(token))
    assert page.status_code == 200
    body = page.json()
    assert body["cards"][0]["code"] == "S95"
    assert body["open_count"] == 1
    assert "secret-token-value" not in page.text


async def test_mizito_token_is_saved_without_being_shown(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    secret = "secret-token-value"
    saved = await client.put("/perf/connections/mizito", headers=auth(token), json={"token": secret})
    assert saved.status_code == 200
    assert secret not in saved.text
    assert saved.json()["saved"] is True
    page = await client.get("/perf/connections", headers=auth(token))
    assert page.status_code == 200
    body = page.json()
    assert body["mizito"]["saved"] is True
    assert body["mizito"]["connected"] is False
    assert "توکن" not in body["mizito"]["reason"]
    assert body["mizito"]["reason"]
    assert secret not in page.text
    assert "sheets" in body
    assert "حساب سرویس" in body["sheets"]["report"]["reason"]
    cleared = await client.put("/perf/connections/mizito", headers=auth(token), json={"token": ""})
    assert cleared.json()["saved"] is False
    again = await client.get("/perf/connections", headers=auth(token))
    assert again.json()["mizito"]["saved"] is False
    assert "توکن میزیتو ذخیره نشده" in again.json()["mizito"]["reason"]


def test_mizito_status_follows_the_latest_probe():
    from app.perf.services.connections import _mizito_status
    from app.perf.services.mizito import mizito_http_reason

    ok = _mizito_status("x" * 12, {"connected": False, "probed_ok": True, "probed_at": "2026-01-01T00:00:00", "probe_reason": ""})
    assert ok["connected"] is True
    bad = _mizito_status("x" * 12, {"connected": True, "probed_ok": False, "probed_at": "2026-01-01T00:00:00", "probe_reason": "توکن میزیتو پذیرفته نشد"})
    assert bad["connected"] is False
    assert "پذیرفته نشد" in bad["reason"]
    assert "۴۰۱" in mizito_http_reason(401, {"msg": "unauthorized"})
    assert "unauthorized" in mizito_http_reason(401, {"msg": "unauthorized"})


def test_google_sheet_error_keeps_status_and_google_text():
    from app.perf.services.connections import google_sheet_reason

    text = google_sheet_reason(403, {"error": {"message": "The caller does not have permission"}})
    assert "۴۰۳" in text
    assert "The caller does not have permission" in text
    assert "توکن" not in text
    missing = google_sheet_reason(404, {})
    assert "۴۰۴" in missing


async def test_campaign_brief_draft_and_save(client, seeded, db):
    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)
    campaign = await client.post(
        "/perf/campaigns",
        headers=headers,
        json={"code": "S96", "title": "شعبه نور", "brand": "ستوده", "goal": "معرفی شعبه نور", "channels": ["sms"]},
    )
    assert campaign.status_code == 201, campaign.text
    campaign_id = campaign.json()["id"]

    draft = await client.post(f"/perf/campaigns/{campaign_id}/brief/draft", headers=headers)
    assert draft.status_code == 200, draft.text
    body = draft.json()
    assert body["source"] in ("llm", "template")
    brief = body["brief"]
    assert brief["slogan"]
    assert len(brief["content_plan"]) == 7
    assert brief["content_plan"][0]["day"] == "شنبه"

    saved = await client.patch(
        f"/perf/campaigns/{campaign_id}",
        headers=headers,
        json={"brief": {"slogan": "نور، جای قرار", "goals": ["معرفی", "جذب ورودی"], "ad_phases": [{"name": "آگاهی", "budget_share": 40}]}},
    )
    assert saved.status_code == 200, saved.text
    saved_brief = saved.json()["brief"]
    assert saved_brief["slogan"] == "نور، جای قرار"
    assert saved_brief["goals"] == ["معرفی", "جذب ورودی"]
    assert saved_brief["ad_phases"][0]["budget_share"] == 40
    assert len(saved_brief["content_plan"]) == 7

    detail = await client.get(f"/perf/campaigns/{campaign_id}", headers=headers)
    assert detail.json()["brief"]["slogan"] == "نور، جای قرار"


def test_isms_credit_probe_reads_balance_and_respects_lockout():
    import asyncio

    import httpx
    from types import SimpleNamespace

    from app.perf.services import connections
    from app.perf.services.connections import probe_isms_credit

    row = SimpleNamespace(organization_id="org-1", code="isms", api_username="user1", api_secret="pass1")

    async def call(handler) -> tuple[dict, dict]:
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            first = await probe_isms_credit(client, row)
            second = await probe_isms_credit(client, row)  # در قفل، دوباره کال نمی‌زند
        return first, second

    def locked_handler(request: httpx.Request) -> httpx.Response:
        connections.calls.append(str(request.url))
        return httpx.Response(200, json={"code": 102, "message": "Login failed 3 times! try ten minutes later"})

    connections.calls = []
    first, second = asyncio.run(call(locked_handler))
    assert first["ok"] is False and "۱۰ دقیقه" in first["reason"]
    assert second["ok"] is False and "قفل" in second["reason"]
    assert len(connections.calls) == 1  # بار دوم از حافظهٔ قفل خواند

    def ok_handler(request: httpx.Request) -> httpx.Response:
        assert request.url.params["username"] == "user1"
        assert request.url.params["password"] == "pass1"
        return httpx.Response(200, json={"code": 0, "credit": 1740096})

    async def ok_call() -> dict:
        connections._ISMS_LOCKED.clear()
        async with httpx.AsyncClient(transport=httpx.MockTransport(ok_handler)) as client:
            return await probe_isms_credit(client, row)

    assert asyncio.run(ok_call())["ok"] is True


def test_panel_probe_reads_balance_live():
    import asyncio
    import json

    import httpx
    from types import SimpleNamespace

    from app.perf.services import connections
    from app.perf.services.connections import probe_panel

    row = SimpleNamespace(code="pishgam", api_base="https://api.example.ir/", api_token="tok-12345678", balance_toman=0)

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.host == "api.example.ir"
        assert request.headers["Authorization"] == "tok-12345678"
        if request.url.path == "/Messages/GetCredit":
            return httpx.Response(200, json={"credit": 17400962, "statusCode": 0})
        if request.url.path == "/Messages/PrivateNumberList":
            return httpx.Response(200, json={"numbers": [{"number": "5000302137682", "type": 1}, {"number": "5000337682", "type": 1}, {"number": "5000302137682", "type": 1}], "statusCode": 0})
        return httpx.Response(404, text="nope")

    def bad_token_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"statusCode": 34})

    def bad_ip_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"statusCode": 35})

    async def main() -> tuple[dict, dict, dict]:
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            out = await probe_panel(client, row)
        async with httpx.AsyncClient(transport=httpx.MockTransport(bad_token_handler)) as client:
            token_bad = await probe_panel(client, row)
        async with httpx.AsyncClient(transport=httpx.MockTransport(bad_ip_handler)) as client:
            ip_bad = await probe_panel(client, row)
        return out, token_bad, ip_bad

    out, token_bad, ip_bad = asyncio.run(main())
    assert out["ok"] is True
    assert out["credit_rial"] == 17400962
    assert row.balance_toman == 17400962 // 10  # موجودی واقعی در دفتر CM نشست
    assert "۱۷٬۴۰۰٬۹۶۲" in out["balance"]
    assert out["lines"] == ["5000302137682", "5000337682"]  # تکراری حذف شده
    assert token_bad["ok"] is False and "توکن" in token_bad["reason"]
    assert ip_bad["ok"] is False and "آی‌پی" in ip_bad["reason"]
    assert json.dumps(out, ensure_ascii=False)  # پاسخ سری‌پذیر است


def test_pishgam_delivery_texts_and_count():
    from app.perf.services.integrations import delivery_text, pishgam_delivered_count

    assert delivery_text(1) == "رسیده به گوشی"
    assert delivery_text("11") == "منقضی"
    assert delivery_text(999) == "999"
    assert pishgam_delivered_count([1, 1, 2, 9, "1", None]) == 3


async def test_pishgam_send_delivery_saves_message_id_and_count(client, seeded, db, monkeypatch):
    import json as jsonlib

    from app.core.config import get_settings
    from app.perf import routes

    await _nadia(db, seeded)
    token = await login(client, "nadia-test@autoco.example")
    headers = auth(token)

    panel = await client.post(
        "/perf/panels",
        headers=headers,
        json={"code": "pishgam", "name": "پیشگام", "sender_line": "5000302137682", "tariff_toman": 100, "balance_toman": 1000, "low_balance_toman": 100},
    )
    assert panel.status_code == 201, panel.text
    panel_id = panel.json()["id"]

    send = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={"panel_id": panel_id, "body": "فونیکس تماس 02147677777", "recipient_count": 2, "pishgam_message_id": 555},
    )
    assert send.status_code == 201, send.text
    send_id = send.json()["id"]
    assert send.json()["pishgam_message_id"] == 555

    async def fake_call(**kwargs):
        assert kwargs["path"] == get_settings().pishgam_report_path
        assert kwargs["method"] == "POST"
        assert kwargs["payload"] == {"messageId": 555, "recipientNumbers": ["09121234567", "09121234568"]}
        return {
            "status_code": 200,
            "body": jsonlib.dumps({"statusCode": 0, "numberDeliveryStatus": [
                {"toNumber": "09121234567", "deliveryStatus": 1},
                {"toNumber": "09121234568", "deliveryStatus": 2},
            ]}),
            "url": "fake",
        }

    monkeypatch.setattr(routes, "pishgam_call", fake_call)
    result = await client.post(
        f"/perf/sms/sends/{send_id}/pishgam-delivery",
        headers=headers,
        json={"recipients": ["09121234567", "09121234568"]},
    )
    assert result.status_code == 200, result.text
    data = result.json()
    assert data["ok"] is True
    assert data["delivered_count"] == 1
    assert data["rows"][0]["text"] == "رسیده به گوشی"
    assert data["rows"][1]["text"] == "نرسیده به گوشی"

    sends = await client.get("/perf/sms/sends", headers=headers)
    row = next(item for item in sends.json() if item["id"] == send_id)
    assert row["delivered_count"] == 1
    assert row["pishgam_message_id"] == 555

    # بدون شناسهٔ ذخیره‌شده، خطای گویا
    other = await client.post(
        "/perf/sms/sends",
        headers=headers,
        json={"panel_id": panel_id, "body": "فونیکس تماس 02147677777", "recipient_count": 2},
    )
    other_id = other.json()["id"]
    missing = await client.post(
        f"/perf/sms/sends/{other_id}/pishgam-delivery",
        headers=headers,
        json={"recipients": ["09121234567"]},
    )
    assert missing.status_code == 422
    assert "شناسهٔ" in missing.json()["detail"]


def test_normalize_brief_fills_the_week_and_keeps_shapes():
    from types import SimpleNamespace

    from app.perf.services.brief import normalize_brief

    campaign = SimpleNamespace(brand="ستوده", title="شعبه نور", goal="", budget_toman=None)
    out = normalize_brief({"goals": ["یک"], "content_plan": [{"day": "شنبه", "details": "معرفی"}], "ad_phases": [{"name": "آگاهی", "budget_share": "140"}], "hashtags": None}, campaign)
    assert out["goals"] == ["یک"]
    assert [row["day"] for row in out["content_plan"]] == ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه", "جمعه"]
    assert out["ad_phases"][0]["budget_share"] == 100
    assert out["hashtags"] == []
