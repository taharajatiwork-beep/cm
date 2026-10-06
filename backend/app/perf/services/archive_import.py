"""کمپین‌های دفتر شماره را به ردیف پرداخت، پیامک و گزارش کانال وصل می‌کند."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.perf.campaign_register import LEGACY, NUMBERS, tag_brand
from app.perf.models import PmCampaign, PmCost, PmPage, PmSmsSend, PmStory

ASSETS = Path(__file__).resolve().parents[4] / "assets"
_REPORTER_SHEETS = (
    ("Instagram-S", "story"),
    ("Instagram-R", "story"),
    ("BaleRubika", "post"),
    ("Native Ads", "banner"),
    ("Native Video Ads", "banner"),
    ("Pre-roll Video Ads", "banner"),
    ("Divar", "banner"),
    ("App", "banner"),
)


def known_campaigns() -> dict[str, dict]:
    book: dict[str, dict] = {}
    for row in [*NUMBERS, *LEGACY]:
        book.setdefault(row["code"], row)
    return book


def _money(value, rial: bool = False) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        amount = int(value)
    else:
        text = str(value).replace(",", "").replace("،", "").strip()
        try:
            amount = int(float(text))
        except ValueError:
            return 0
    if rial and amount:
        return amount // 10
    return amount


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _when(value):
    if value is None or value == "":
        return None
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return value
    raw = str(value).strip().replace("/", "-")[:10]
    try:
        from datetime import date

        return date.fromisoformat(raw)
    except ValueError:
        return None


def _header(row) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _text(cell).lower()
        if key and key not in found:
            found[key] = index
    return found


def _pick(row, header: dict[str, int], *needles: str):
    for needle in needles:
        for key, index in header.items():
            if needle in key and index < len(row):
                return row[index]
    return None


def _rows(path: Path, sheet: str):
    book = load_workbook(path, data_only=True, read_only=True)
    if sheet not in book.sheetnames:
        return
    ws = book[sheet]
    header = None
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        if header is None:
            if any(_text(cell) for cell in cells):
                header = _header(cells)
            continue
        if any(cell is not None and _text(cell) for cell in cells):
            yield header, cells


async def import_sotoodeh(db: AsyncSession, organization_id: UUID) -> dict:
    book = known_campaigns()
    campaigns = {
        row.code: row
        for row in (
            await db.scalars(select(PmCampaign).where(PmCampaign.organization_id == organization_id))
        ).all()
    }
    made = 0
    for code, row in book.items():
        current = campaigns.get(code)
        title = row.get("title") or code
        brand = tag_brand(row.get("tag") or "")
        if current is None:
            current = PmCampaign(
                organization_id=organization_id,
                code=code,
                title=title,
                brand=brand,
                status="done",
                channels="sms,ig_story",
            )
            db.add(current)
            campaigns[code] = current
            made += 1
        else:
            if title and title != code:
                current.title = title
            if brand:
                current.brand = brand
    await db.flush()

    costs = await _import_log(db, organization_id, campaigns)
    sends = await _import_sms(db, organization_id, campaigns)
    stories = await _import_reporter(db, organization_id, campaigns)
    await db.flush()
    return {"campaigns": made, "costs": costs, "sends": sends, "stories": stories}


async def _import_log(db: AsyncSession, organization_id: UUID, campaigns: dict[str, PmCampaign]) -> int:
    path = next(ASSETS.glob("*Cost Tracking.xlsx"), None)
    if path is None:
        return 0
    stored = list((await db.scalars(select(PmCost).where(PmCost.organization_id == organization_id))).all())
    number = max((row.cost_number or 0 for row in stored), default=0)
    existing = {
        (
            (row.campaign_code or "").strip(),
            row.amount_toman or 0,
            row.spent_on.isoformat() if row.spent_on else "",
            (row.description or "")[:80],
        )
        for row in stored
    }
    added = 0
    for header, cells in _rows(path, "Log"):
        code = _text(_pick(cells, header, "کمپین"))
        campaign = campaigns.get(code)
        if campaign is None:
            continue
        amount = _money(_pick(cells, header, "مبلغ"))
        if amount <= 0:
            continue
        spent = _when(_pick(cells, header, "تاریخ"))
        description = _text(_pick(cells, header, "شرح"))[:400]
        key = (code, amount, spent.isoformat() if spent else "", description[:80])
        if key in existing:
            continue
        existing.add(key)
        number += 1
        db.add(
            PmCost(
                organization_id=organization_id,
                cost_number=number,
                spent_on=spent,
                amount_toman=amount,
                kind=_text(_pick(cells, header, "نوع"))[:80],
                detail=_text(_pick(cells, header, "تفضیلی"))[:160],
                brand=campaign.brand or _text(_pick(cells, header, "شرکت"))[:80],
                campaign_code=code,
                cost_center=_text(_pick(cells, header, "مرکز"))[:120],
                description=description,
                counts_for_cpl=True,
            )
        )
        added += 1
    return added


async def _import_sms(db: AsyncSession, organization_id: UUID, campaigns: dict[str, PmCampaign]) -> int:
    path = ASSETS / "ارسال پیامک.xlsx"
    if not path.exists():
        return 0
    existing = {
        (row.campaign_id, (row.body or "")[:120], row.recipient_count or 0)
        for row in (await db.scalars(select(PmSmsSend).where(PmSmsSend.organization_id == organization_id))).all()
    }
    added = 0
    for sheet in ("Sheet1", "Sheet2"):
        for header, cells in _rows(path, sheet):
            code = _text(_pick(cells, header, "کمپین"))
            campaign = campaigns.get(code)
            if campaign is None:
                continue
            body = _text(_pick(cells, header, "شرح", "متن"))
            count = _money(_pick(cells, header, "گیرنده", "تعداد"))
            if not body and count <= 0:
                continue
            key = (campaign.id, body[:120], count)
            if key in existing:
                continue
            existing.add(key)
            spent = _when(_pick(cells, header, "تاریخ"))
            cost_cell = None
            rial = False
            for name, index in header.items():
                if "هزینه" in name or name == "sum":
                    cost_cell = cells[index] if index < len(cells) else None
                    rial = "ریال" in name
                    break
            db.add(
                PmSmsSend(
                    organization_id=organization_id,
                    campaign_id=campaign.id,
                    body=body,
                    pages=1,
                    recipient_count=count,
                    sent_at=datetime(spent.year, spent.month, spent.day, tzinfo=UTC) if spent else None,
                    cost_toman=_money(cost_cell, rial),
                )
            )
            added += 1
    return added


async def _import_reporter(db: AsyncSession, organization_id: UUID, campaigns: dict[str, PmCampaign]) -> int:
    path = ASSETS / "reporter-1405.xlsx"
    if not path.exists():
        return 0
    pages = {
        row.name: row
        for row in (await db.scalars(select(PmPage).where(PmPage.organization_id == organization_id))).all()
    }
    existing = {
        (row.campaign_id, row.page_id, row.views or 0, row.cost_toman or 0, row.leads or 0)
        for row in (await db.scalars(select(PmStory).where(PmStory.organization_id == organization_id))).all()
    }
    added = 0
    book = load_workbook(path, data_only=True, read_only=True)
    for sheet, kind in _REPORTER_SHEETS:
        if sheet not in book.sheetnames:
            continue
        header = None
        current = ""
        for raw in book[sheet].iter_rows(values_only=True):
            cells = list(raw)
            if header is None:
                labels = [_text(cell).lower() for cell in cells]
                if any("campaign" in label or "کمپین" in label for label in labels):
                    header = _header(cells)
                continue
            code = _text(_pick(cells, header, "campaign", "کمپین"))
            if code in campaigns:
                current = code
            elif code:
                current = ""
            campaign = campaigns.get(current)
            if campaign is None:
                continue
            page_name = _text(_pick(cells, header, "page"))
            views = _money(_pick(cells, header, "view"))
            leads = _money(_pick(cells, header, "lead"))
            cost = _money(_pick(cells, header, "cost"))
            if not page_name and views <= 0 and leads <= 0 and cost <= 0:
                continue
            page = pages.get(page_name) if page_name else None
            if page_name and page is None:
                page = PmPage(organization_id=organization_id, name=page_name[:160])
                db.add(page)
                await db.flush()
                pages[page_name] = page
            published = _when(_pick(cells, header, "date"))
            key = (campaign.id, page.id if page else None, views, cost, leads)
            if key in existing:
                continue
            existing.add(key)
            db.add(
                PmStory(
                    organization_id=organization_id,
                    campaign_id=campaign.id,
                    page_id=page.id if page else None,
                    kind=kind,
                    status="published",
                    published_on=published,
                    cost_toman=cost,
                    views=views,
                    leads=leads,
                )
            )
            added += 1
    return added
