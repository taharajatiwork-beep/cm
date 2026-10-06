"""خروجی و ورود اکسل Cost Tracking و reporter."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

COST_HEADERS = [
    "شماره هزینه",
    "تاریخ",
    "مبلغ (تومان)",
    "نوع",
    "تفضیلی",
    "برند",
    "کمپین",
    "مرکز هزینه",
    "شرح هزینه",
]


def _book() -> tuple[Workbook, BytesIO]:
    return Workbook(), BytesIO()


def _save(book: Workbook) -> bytes:
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def costs_xlsx(rows: list[dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Cost Tracking"
    sheet.append(COST_HEADERS)
    for row in rows:
        sheet.append(
            [
                row.get("cost_number"),
                _cell_date(row.get("spent_on")),
                row.get("amount_toman"),
                row.get("kind"),
                row.get("detail"),
                row.get("brand"),
                row.get("campaign_code"),
                row.get("cost_center"),
                row.get("description"),
            ]
        )
    return _save(book)


def reporter_xlsx(report_rows: list[dict], sms_rows: list[dict], story_rows: list[dict]) -> bytes:
    book = Workbook()
    total = book.active
    total.title = "Total"
    total.append(["Platform", "Ad Channel", "Cost", "Lead", "CPL"])
    for row in report_rows:
        for channel in row.get("channels") or []:
            leads = channel.get("leads") or 0
            cost = channel.get("cost") or 0
            total.append(
                [
                    "Performance",
                    channel.get("label"),
                    cost,
                    leads,
                    round(cost / leads, 1) if leads else None,
                ]
            )
        total.append(["Performance", row.get("code"), row.get("cost_toman"), row.get("leads"), row.get("cpl")])

    campaign = book.create_sheet("Campaign")
    campaign.append(["Campaign", "Channel", "Cost", "View", "Lead", "CPL", "CPV"])
    for row in report_rows:
        for channel in row.get("channels") or []:
            campaign.append(
                [
                    row.get("code"),
                    channel.get("label"),
                    channel.get("cost"),
                    channel.get("views"),
                    channel.get("leads"),
                    round(channel["cost"] / channel["leads"], 1) if channel.get("leads") else None,
                    round(channel["cost"] / channel["views"], 1) if channel.get("views") else None,
                ]
            )

    sms = book.create_sheet("SMS")
    sms.append(["کمپین", "تاریخ", "خط", "گیرنده", "صفحه", "هزینه", "تحویل", "تماس", "ورودی", "متن"])
    for row in sms_rows:
        sms.append(
            [
                row.get("campaign_code"),
                _cell_date(row.get("sent_at")),
                row.get("sender_line"),
                row.get("recipient_count"),
                row.get("pages"),
                row.get("cost_toman"),
                row.get("delivered_count"),
                row.get("calls"),
                row.get("leads"),
                row.get("body"),
            ]
        )

    story = book.create_sheet("Instagram-S")
    story.append(
        ["Campaign", "Page Name", "Publication Date", "Cost", "Views", "Reach", "Link Clicks", "Shares", "Replies", "Exited", "Lead", "CPR", "CPV", "CPC", "CTR", "CPL"]
    )
    for row in story_rows:
        story.append(
            [
                row.get("campaign_code"),
                row.get("page"),
                _cell_date(row.get("published_on")),
                row.get("cost_toman"),
                row.get("views"),
                row.get("reach"),
                row.get("link_clicks"),
                row.get("shares"),
                row.get("replies"),
                row.get("exited"),
                row.get("leads"),
                row.get("cpr"),
                row.get("cpv"),
                row.get("cpc"),
                row.get("ctr"),
                row.get("cpl"),
            ]
        )
    return _save(book)


def blank_template(kind: str) -> bytes:
    book = Workbook()
    sheet = book.active
    if kind == "leads":
        sheet.title = "Leads"
        sheet.append(["شماره", "تاریخ", "کمپین", "کانال", "مرحله"])
    else:
        sheet.title = "Cost Tracking"
        sheet.append(COST_HEADERS)
    return _save(book)


def _cell_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "")


def _header_map(header: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    aliases = {
        "cost_number": ("شماره هزینه", "شماره"),
        "spent_on": ("تاریخ",),
        "amount": ("مبلغ",),
        "kind": ("نوع",),
        "detail": ("تفضیلی",),
        "brand": ("برند",),
        "campaign": ("کمپین",),
        "center": ("مرکز",),
        "description": ("شرح",),
        "phone": ("شماره", "موبایل", "تلفن", "phone"),
        "channel": ("کانال", "channel"),
        "stage": ("مرحله", "stage"),
    }
    for key, names in aliases.items():
        for i, cell in enumerate(header):
            if any(name in cell for name in names):
                found[key] = i
                break
    return found


_INSIGHT_NAMES = {
    "views": ("نمایش", "view", "impress"),
    "reach": ("دسترس", "reach"),
    "replies": ("پاسخ", "reply", "comment"),
    "shares": ("اشتراک", "share"),
    "link_clicks": ("کلیک", "click"),
    "exited": ("خروج", "exit"),
    "leads": ("ورودی", "لید", "lead"),
}


def parse_insight_sheet(data: bytes) -> dict[str, int]:
    """عدد اینسایت از ستون‌های ثابت اکسل، بدون مدل زبانی."""
    header, rows = read_table(data)
    found: dict[str, int] = {}
    for key, names in _INSIGHT_NAMES.items():
        for index, cell in enumerate(header):
            if any(name in cell for name in names):
                found[key] = index
                break
    totals: dict[str, int] = {}
    for line in rows:
        if not any(line):
            continue
        for key, index in found.items():
            if index >= len(line) or line[index] in (None, ""):
                continue
            try:
                totals[key] = totals.get(key, 0) + int(float(line[index]))
            except (TypeError, ValueError):
                continue
    return totals


def read_table(data: bytes) -> tuple[list[str], list[list]]:
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = book.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    if not rows:
        return [], []
    header = [str(cell or "").strip() for cell in rows[0]]
    return header, rows[1:]


def parse_cost_rows(data: bytes) -> list[dict]:
    header, rows = read_table(data)
    idx = _header_map(header)
    parsed = []
    for row in rows:
        if not any(row):
            continue

        def cell(key: str):
            i = idx.get(key)
            if i is None or i >= len(row):
                return None
            return row[i]

        amount = cell("amount")
        try:
            amount_n = int(float(amount or 0))
        except (TypeError, ValueError):
            amount_n = 0
        number = cell("cost_number")
        try:
            number_n = int(number) if number not in (None, "") else 0
        except (TypeError, ValueError):
            number_n = 0
        parsed.append(
            {
                "cost_number": number_n,
                "spent_on": str(cell("spent_on") or "")[:10],
                "amount_toman": amount_n,
                "kind": str(cell("kind") or ""),
                "detail": str(cell("detail") or ""),
                "brand": str(cell("brand") or ""),
                "campaign_code": str(cell("campaign") or ""),
                "cost_center": str(cell("center") or ""),
                "description": str(cell("description") or ""),
            }
        )
    return parsed


def parse_lead_rows(data: bytes) -> list[dict]:
    header, rows = read_table(data)
    idx = _header_map(header)
    parsed = []
    for row in rows:
        if not any(row):
            continue

        def cell(key: str):
            i = idx.get(key)
            if i is None or i >= len(row):
                return ""
            return row[i]

        parsed.append(
            {
                "phone": str(cell("phone") or ""),
                "occurred_on": str(cell("spent_on") or "")[:10],
                "campaign_code": str(cell("campaign") or ""),
                "channel": str(cell("channel") or ""),
                "stage": str(cell("stage") or ""),
            }
        )
    return parsed
