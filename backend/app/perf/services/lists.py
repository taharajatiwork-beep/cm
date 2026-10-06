"""یکدست‌سازی شماره و آماده‌سازی لیست دیدار."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

_DIGIT = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
_HEADER_KEYS = ("موبایل", "تلفن", "شماره", "phone", "mobile", "cellphone")


def normalize_phone(raw: object) -> str | None:
    text = str(raw or "").translate(_DIGIT)
    digits = "".join(ch for ch in text if ch.isdigit())
    if digits.startswith("0098"):
        digits = digits[2:]
    if digits.startswith("98") and len(digits) == 12:
        national = digits
    elif digits.startswith("0") and len(digits) == 11 and digits[1] == "9":
        national = "98" + digits[1:]
    elif len(digits) == 10 and digits.startswith("9"):
        national = "98" + digits
    else:
        return None
    if not national.startswith("989") or len(national) != 12:
        return None
    return national


def panel_number(national: str) -> str:
    """شماره برای آپلود پنل: بدون صفر اول."""
    return national[2:] if national.startswith("98") else national


def clean_numbers(raw_phones: list[str], *, blacklist: set[str], recent: set[str]) -> dict:
    seen: set[str] = set()
    final: list[str] = []
    duplicate = blacklist_n = recent_n = invalid = 0
    for raw in raw_phones:
        phone = normalize_phone(raw)
        if not phone:
            invalid += 1
            continue
        if phone in seen:
            duplicate += 1
            continue
        seen.add(phone)
        if phone in blacklist:
            blacklist_n += 1
            continue
        if phone in recent:
            recent_n += 1
            continue
        final.append(phone)
    return {
        "raw_count": len(raw_phones),
        "duplicate_count": duplicate,
        "blacklist_count": blacklist_n,
        "recent_count": recent_n,
        "invalid_count": invalid,
        "final_count": len(final),
        "phones": final,
    }


def phones_from_xlsx(data: bytes) -> list[str]:
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = book.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    if not rows:
        return []
    header = [str(cell or "").strip() for cell in rows[0]]
    index = 0
    has_header = False
    for i, name in enumerate(header):
        lowered = name.lower()
        if any(key in name or key in lowered for key in _HEADER_KEYS):
            index = i
            has_header = True
            break
    start = 1 if has_header else 0
    if not has_header and normalize_phone(header[0] if header else "") is None and any(header):
        # ردیف اول عنوان ناشناخته است؛ از ستون اول داده می‌خوانیم
        start = 1
    found: list[str] = []
    for row in rows[start:]:
        if index >= len(row) or row[index] is None:
            continue
        found.append(str(row[index]))
    return found


def workbook_chunks(phones: list[str], chunk_size: int) -> bytes:
    size = max(1, chunk_size or 15000)
    book = Workbook()
    first = book.active
    chunks = [phones[offset : offset + size] for offset in range(0, len(phones), size)] or [[]]
    for index, chunk in enumerate(chunks):
        sheet = first if index == 0 else book.create_sheet()
        sheet.title = f"part-{index + 1}"
        sheet.append(["mobile"])
        for phone in chunk:
            sheet.append([panel_number(phone)])
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()
