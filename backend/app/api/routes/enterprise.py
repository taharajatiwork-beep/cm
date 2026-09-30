"""Admin, audit, approvals, reports, retention, and integration adapters."""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime
from uuid import UUID

from fastapi import APIRouter, Depends, Header, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import AuthContext, get_current_user, handle_authz
from app.api.pagination import encode_cursor, envelope
from app.core.jalali import format_jalali
from app.core.permissions import AuthorizationError, PermissionCode, require_permission
from app.db.session import get_db
from app.models import Conversation, Permission, Role, RolePermission, User
from app.models.enterprise import ApprovalRequest, ApprovalStep, IdempotencyRecord, KpiSnapshot, ReportSchedule
from app.services import write_audit
from app.services.enterprise_ops import EnterpriseService, notify_current_approver

router = APIRouter(tags=["enterprise"])


class ApprovalIn(BaseModel):
    kind: str
    title: str
    payload: dict = Field(default_factory=dict)
    approver_ids: list[UUID]


class DecisionIn(BaseModel):
    decision: str
    comment: str | None = None


class RoleUpdate(BaseModel):
    permission_codes: list[str]
    data_scope_level: str | None = None


class TeamMembershipIn(BaseModel):
    user_id: UUID
    team_id: UUID


class UserAdminUpdate(BaseModel):
    """Org-admin actions on a single user. Both fields are optional."""

    is_active: bool | None = None
    password: str | None = None


class ReportScheduleIn(BaseModel):
    name: str
    channel: str = "email"
    hour_utc: int = 5


def _admin(ctx: AuthContext) -> None:
    try:
        require_permission(ctx.permissions, PermissionCode.ORG_ADMIN)
    except AuthorizationError as e:
        raise handle_authz(e) from e


@router.get("/audit")
async def list_audit(
    actor_id: UUID | None = None,
    entity_type: str | None = None,
    action: str | None = None,
    day: str | None = None,
    limit: int = 50,
    envelope_mode: bool = False,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    try:
        rows = await EnterpriseService(db).list_audit(
            ctx.organization_id,
            actor_permissions=ctx.permissions,
            actor_id=actor_id,
            entity_type=entity_type,
            action=action,
            day=day,
            limit=limit,
        )
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    except AuthorizationError as e:
        raise handle_authz(e) from e
    items = [
        {
            "id": str(r.id),
            "action": r.action,
            "entity_type": r.entity_type,
            "entity_id": str(r.entity_id),
            "actor_id": str(r.actor_id) if r.actor_id else None,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "created_at_jalali": format_jalali(r.created_at) if r.created_at else None,
        }
        for r in rows
    ]
    if envelope_mode:
        last_id = items[-1].get("id") if items else None
        nxt = encode_cursor(last_id) if (len(items) == limit and last_id) else None
        return envelope(items, total=len(items), next_cursor=nxt)
    return items


@router.get("/audit.xlsx")
async def audit_xlsx(
    actor_id: UUID | None = None,
    entity_type: str | None = None,
    action: str | None = None,
    day: str | None = None,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """خروجی اکسل ممیزی — همان فیلترهای تب ممیزی."""
    from io import BytesIO

    from openpyxl import Workbook

    _admin(ctx)
    try:
        rows = await EnterpriseService(db).list_audit(
            ctx.organization_id,
            actor_permissions=ctx.permissions,
            actor_id=actor_id,
            entity_type=entity_type,
            action=action,
            day=day,
            limit=1000,
        )
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    wb = Workbook()
    ws = wb.active
    if ws is None:  # openpyxl types `active` as optional
        ws = wb.create_sheet()
    ws.title = "audit"
    ws.append(["action", "entity_type", "entity_id", "actor_id", "created_at"])
    for r in rows:
        ws.append(
            [
                r.action,
                r.entity_type,
                str(r.entity_id),
                str(r.actor_id) if r.actor_id else "",
                r.created_at.isoformat() if r.created_at else "",
            ]
        )
    buffer = BytesIO()
    wb.save(buffer)
    await write_audit(
        db,
        organization_id=ctx.organization_id,
        actor_id=ctx.user.id,
        action="report.export",
        entity_type="report",
        entity_id=ctx.organization_id,
        details={"format": "xlsx", "kind": "audit"},
    )
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=audit.xlsx"},
    )


@router.post("/approvals")
async def create_approval(
    body: ApprovalIn,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
):
    if idempotency_key:
        existing = await db.scalar(
            select(IdempotencyRecord).where(
                IdempotencyRecord.organization_id == ctx.organization_id,
                IdempotencyRecord.key == idempotency_key,
            )
        )
        if existing:
            return json.loads(existing.response_body)
    try:
        row = await EnterpriseService(db).request_approval(
            organization_id=ctx.organization_id,
            requester_id=ctx.user.id,
            kind=body.kind,
            title=body.title,
            payload=body.payload,
            approver_ids=body.approver_ids,
        )
    except ValueError as e:
        # خود-تأییدی، تأییدکنندهٔ نامعتبر و سایر خطاهای دامنه → ۴۰۰ نه ۵۰۰
        raise HTTPException(400, str(e)) from e
    await notify_current_approver(db, row)
    payload = {"id": str(row.id), "status": row.status, "current_step": row.current_step}
    if idempotency_key:
        db.add(
            IdempotencyRecord(
                organization_id=ctx.organization_id,
                user_id=ctx.user.id,
                key=idempotency_key,
                response_status=200,
                response_body=json.dumps(payload),
            )
        )
        await db.flush()
    return payload


@router.post("/approvals/{request_id}/decision")
async def decide_approval(
    request_id: UUID,
    body: DecisionIn,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if body.decision not in {"approved", "rejected"}:
        raise HTTPException(400, "decision must be approved or rejected")
    try:
        row = await EnterpriseService(db).decide(
            organization_id=ctx.organization_id,
            request_id=request_id,
            actor_id=ctx.user.id,
            decision=body.decision,
            comment=body.comment,
        )
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    if row.status == "pending" and body.decision == "approved":
        await notify_current_approver(db, row)
    return {"id": str(row.id), "status": row.status, "current_step": row.current_step}


@router.get("/approvals")
async def list_approvals(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    rows = list(
        (
            await db.execute(
                select(ApprovalRequest)
                .where(ApprovalRequest.organization_id == ctx.organization_id)
                .order_by(ApprovalRequest.created_at.desc())
            )
        ).scalars()
    )
    out = []
    for row in rows:
        steps = list(
            (
                await db.execute(
                    select(ApprovalStep)
                    .where(ApprovalStep.request_id == row.id)
                    .order_by(ApprovalStep.step_no)
                )
            ).scalars()
        )
        out.append(
            {
                "id": str(row.id),
                "title": row.title,
                "kind": row.kind,
                "status": row.status,
                "current_step": row.current_step,
                "requester_id": str(row.requester_id),
                "created_at": row.created_at.isoformat() if row.created_at else None,
                "steps": [
                    {
                        "step_no": step.step_no,
                        "approver_id": str(step.approver_id),
                        "decision": step.decision,
                        "comment": step.comment,
                    }
                    for step in steps
                ],
            }
        )
    return out


@router.get("/reports/kpis")
async def kpis(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    try:
        require_permission(ctx.permissions, PermissionCode.EVENT_READ)
    except AuthorizationError as e:
        raise handle_authz(e) from e
    data = await EnterpriseService(db).kpis(ctx.organization_id)
    data["as_of_jalali"] = format_jalali(datetime.now(UTC))
    return data


@router.get("/reports/kpis.csv")
async def kpis_csv(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    try:
        require_permission(ctx.permissions, PermissionCode.EVENT_READ)
    except AuthorizationError as e:
        raise handle_authz(e) from e
    data = await _kpi_payload(ctx, db)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["metric", "value"])
    for key, value in data.items():
        writer.writerow([key, value])
    await write_audit(
        db,
        organization_id=ctx.organization_id,
        actor_id=ctx.user.id,
        action="report.export",
        entity_type="report",
        entity_id=ctx.organization_id,
        details={"format": "csv"},
    )
    return Response(
        content=buf.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=kpis.csv"},
    )


async def _kpi_payload(ctx: AuthContext, db: AsyncSession) -> dict:
    """Shared KPI fetch behind EVENT_READ — every export format uses this."""
    try:
        require_permission(ctx.permissions, PermissionCode.EVENT_READ)
    except AuthorizationError as e:
        raise handle_authz(e) from e
    data = await EnterpriseService(db).kpis(ctx.organization_id)
    data["as_of_jalali"] = format_jalali(datetime.now(UTC))
    return data


async def _write_export_audit(
    db: AsyncSession, ctx: AuthContext, fmt: str
) -> None:
    await write_audit(
        db,
        organization_id=ctx.organization_id,
        actor_id=ctx.user.id,
        action="report.export",
        entity_type="report",
        entity_id=ctx.organization_id,
        details={"format": fmt},
    )


@router.get("/reports/kpis.xlsx")
async def kpis_xlsx(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    from io import BytesIO

    from openpyxl import Workbook

    data = await _kpi_payload(ctx, db)
    wb = Workbook()
    ws = wb.active
    if ws is None:  # openpyxl types `active` as optional
        ws = wb.create_sheet()
    ws.title = "KPIs"
    ws.append(["metric", "value"])
    for key, value in data.items():
        ws.append([key, value])
    buffer = BytesIO()
    wb.save(buffer)
    await _write_export_audit(db, ctx, "xlsx")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=kpis.xlsx"},
    )


@router.get("/reports/kpis.pdf")
async def kpis_pdf(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    from app.core.pdf import build_pdf

    data = await _kpi_payload(ctx, db)
    lines = [f"{key}: {value}" for key, value in data.items()]
    payload = build_pdf(
        title="CM KPI Report",
        lines=[f"as_of (jalali): {data.get('as_of_jalali', '')}", *lines],
    )
    await _write_export_audit(db, ctx, "pdf")
    return Response(
        content=payload,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=kpis.pdf"},
    )


@router.get("/admin/overview")
async def admin_overview(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    """نمای کلی سازمان — فقط برای دارندگان org.admin."""
    _admin(ctx)
    data = await EnterpriseService(db).admin_overview(ctx.organization_id)
    data["as_of_jalali"] = format_jalali(datetime.now(UTC))
    return data


@router.post("/reports/schedules")
async def schedule_report(
    body: ReportScheduleIn,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _admin(ctx)
    row = await EnterpriseService(db).schedule_report(
        ctx.organization_id, body.name, body.channel, body.hour_utc
    )
    return {"id": str(row.id), "name": row.name, "channel": row.channel}


@router.get("/reports/schedules")
async def list_schedules(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    _admin(ctx)
    rows = list(
        (
            await db.execute(
                select(ReportSchedule).where(ReportSchedule.organization_id == ctx.organization_id)
            )
        ).scalars()
    )
    return [
        {"id": str(row.id), "name": row.name, "channel": row.channel, "hour_utc": row.hour_utc, "enabled": row.enabled}
        for row in rows
    ]


@router.post("/reports/kpis/snapshot")
async def snapshot_kpis(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    try:
        require_permission(ctx.permissions, PermissionCode.EVENT_READ)
    except AuthorizationError as e:
        raise handle_authz(e) from e
    data = await EnterpriseService(db).kpis(ctx.organization_id)
    db.add(KpiSnapshot(organization_id=ctx.organization_id, payload_json=json.dumps(data)))
    await db.flush()
    return data


@router.get("/reports/kpis/history")
async def kpi_history(ctx: AuthContext = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    try:
        require_permission(ctx.permissions, PermissionCode.EVENT_READ)
    except AuthorizationError as e:
        raise handle_authz(e) from e
    rows = list(
        (
            await db.execute(
                select(KpiSnapshot)
                .where(KpiSnapshot.organization_id == ctx.organization_id)
                .order_by(KpiSnapshot.created_at.desc())
                .limit(30)
            )
        ).scalars()
    )
    return [
        {
            "id": str(row.id),
            "created_at": row.created_at.isoformat() if row.created_at else None,
            "created_at_jalali": format_jalali(row.created_at) if row.created_at else None,
            "kpis": json.loads(row.payload_json),
        }
        for row in rows
    ]


@router.post("/org/users/{user_id}/teams")
async def add_user_team(
    user_id: UUID,
    body: TeamMembershipIn,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _admin(ctx)
    if body.user_id != user_id:
        raise HTTPException(400, "user_id mismatch")
    try:
        row = await EnterpriseService(db).add_team_membership(ctx.organization_id, user_id, body.team_id)
    except ValueError as e:
        raise HTTPException(404, str(e)) from e
    await write_audit(
        db,
        organization_id=ctx.organization_id,
        actor_id=ctx.user.id,
        action="user.team_added",
        entity_type="user",
        entity_id=user_id,
        details={"team_id": str(body.team_id)},
    )
    return {"id": str(row.id), "team_id": str(row.team_id)}


@router.patch("/org/users/{user_id}")
async def admin_update_user(
    user_id: UUID,
    body: UserAdminUpdate,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Deactivate a user or reset their password. Org admins only."""
    from app.core.security import hash_password, validate_password

    _admin(ctx)
    user = await db.get(User, user_id)
    if not user or user.organization_id != ctx.organization_id:
        raise HTTPException(404, "User not found")

    if body.is_active is not None and body.is_active != user.is_active:
        user.is_active = body.is_active
        await write_audit(
            db,
            organization_id=ctx.organization_id,
            actor_id=ctx.user.id,
            action="user.deactivated" if not body.is_active else "user.activated",
            entity_type="user",
            entity_id=user.id,
            details={"is_active": body.is_active},
        )

    if body.password is not None:
        problem = validate_password(body.password)
        if problem:
            raise HTTPException(400, problem)
        user.hashed_password = hash_password(body.password)
        user.failed_login_count = 0
        user.locked_until = None
        await write_audit(
            db,
            organization_id=ctx.organization_id,
            actor_id=ctx.user.id,
            action="user.password_reset",
            entity_type="user",
            entity_id=user.id,
            details={},
        )

    await db.flush()
    return {"id": str(user.id), "is_active": user.is_active}


@router.get("/org/roles/{role_id}/permissions")
async def get_role_permissions(
    role_id: UUID,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """The permissions and data scope a role currently holds, for the admin UI."""
    from app.models import DataScope

    _admin(ctx)
    role = await db.get(Role, role_id)
    if not role or role.organization_id != ctx.organization_id:
        raise HTTPException(404, "Role not found")
    rows = list(
        (
            await db.execute(
                select(RolePermission).where(RolePermission.role_id == role.id)
            )
        ).scalars()
    )
    codes: list[str] = []
    for row in rows:
        perm = await db.get(Permission, row.permission_id)
        if perm:
            codes.append(perm.code)
    scope = await db.scalar(select(DataScope).where(DataScope.role_id == role.id))
    return {
        "id": str(role.id),
        "code": role.code,
        "name": role.name,
        "permission_codes": sorted(codes),
        "data_scope_level": scope.level if scope else None,
    }


@router.patch("/org/roles/{role_id}")
async def update_role(
    role_id: UUID,
    body: RoleUpdate,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _admin(ctx)
    role = await db.get(Role, role_id)
    if not role or role.organization_id != ctx.organization_id:
        raise HTTPException(404, "Role not found")
    from app.models import DataScope

    existing = list(
        (await db.execute(select(RolePermission).where(RolePermission.role_id == role.id))).scalars()
    )
    for row in existing:
        await db.delete(row)
    for code in body.permission_codes:
        perm = await db.scalar(select(Permission).where(Permission.code == code))
        if perm:
            db.add(RolePermission(role_id=role.id, permission_id=perm.id))
    if body.data_scope_level:
        scope = await db.scalar(select(DataScope).where(DataScope.role_id == role.id))
        if scope:
            scope.level = body.data_scope_level
        else:
            db.add(DataScope(role_id=role.id, level=body.data_scope_level))
    await write_audit(
        db,
        organization_id=ctx.organization_id,
        actor_id=ctx.user.id,
        action="role.updated",
        entity_type="role",
        entity_id=role.id,
        details={"permissions": body.permission_codes, "scope": body.data_scope_level},
    )
    await db.flush()
    return {"id": str(role.id), "code": role.code}


@router.post("/admin/retention/archive")
async def archive_old_conversations(
    older_than_days: int = 365,
    ctx: AuthContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _admin(ctx)
    cutoff = datetime.now(UTC).timestamp() - older_than_days * 86400
    rows = list(
        (
            await db.execute(
                select(Conversation).where(Conversation.organization_id == ctx.organization_id)
            )
        ).scalars()
    )
    archived = 0
    for conv in rows:
        created = conv.created_at
        # آرشیو فقط archived_at را پر می‌کند؛ topic دست‌نخورده می‌ماند (Consistency).
        if created and created.timestamp() < cutoff and conv.archived_at is None:
            conv.archived_at = datetime.now(UTC)
            archived += 1
    await write_audit(
        db,
        organization_id=ctx.organization_id,
        actor_id=ctx.user.id,
        action="retention.archive",
        entity_type="organization",
        entity_id=ctx.organization_id,
        details={"archived": archived},
    )
    return {"archived": archived}


@router.get("/integrations/status")
async def integration_status(ctx: AuthContext = Depends(get_current_user)):
    _admin(ctx)
    from app.core.config import get_settings
    from app.integrations.channels import channel_status

    settings = get_settings()
    return channel_status(settings)
